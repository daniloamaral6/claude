"""
Converte o produtos.json exportado da Shopee na planilha-modelo do catálogo Afeturar.

  python3 -I scripts/shopee_para_planilha.py produtos.json saida.xlsx

Decisões (combinadas com a dona da loja):
  - preços iguais aos da Shopee (menor preço = preço base; variações mais caras viram "preço da variação");
  - tudo "Sob encomenda" (o estoque 5 da Shopee é simbólico e é ignorado);
  - descrições sem "impressão 3D" (a marca não se posiciona como impressão 3D);
  - sem fotos nesta etapa; peso/medidas da embalagem vêm do JSON (kg→g, cm→mm).
Nada é inventado: prazo de preparo e medidas da peça ficam vazios.
"""
import json, re, shutil, sys
import openpyxl

CATEGORIAS = {
    "vasos-decorativos-rib": "Vasos",
    "vasos-decorativos-archer-slim-conjunto-moderno-para-sala-e-a": "Vasos",
    "nossa-senhora-aparecida-decorativa": "Fé",
    "placa-decorativa-orai-e-vigiai": "Fé",
    "bandeja-religiosa-ore-e-confie": "Fé",
    "rogai-por-nos-nossa-senhora-aparecida": "Fé",
    "kit-organizador-de-banheiro-porta-absorvente-porta-protetor-": "Organização",
    "organizador-porta-algodao-porta-pinceis-bandeja": "Organização",
    "porta-ampolas-4-compartimentos-com-tampas-individuais": "Beleza & Acessórios",
    "porta-batom-com-chaveiro-estrela-afeturar": "Beleza & Acessórios",
    "porta-escovinha-de-cilios-com-chaveiro-borboleta": "Beleza & Acessórios",
    "porta-saquinho-higienico-pet-para-cachorro-com-alca-dispense": "Pet",
}
CORRECOES = {"marmore": "Mármore", "pqueno": "Pequeno", "verde menta": "Verde menta", "verde bege": "Verde/Bege", "verde oliva": "Verde oliva"}


def corrigir(v: str) -> str:
    v = v.strip()
    if v.lower() in CORRECOES:
        return CORRECOES[v.lower()]
    partes = re.split(r"(/)", v)
    return "".join(CORRECOES.get(p.lower(), p) for p in partes)


def sem_3d(t: str) -> str:
    t = t.replace("🖨️ ", "").replace("🖨️", "")
    t = re.sub(r"características d[ao]s? impressão 3D", "características do processo de fabricação", t)
    t = re.sub(r"(?i)\b(em|por|através de|por meio de) impressão 3D", "sob demanda", t)
    t = re.sub(r"(?i)impressão 3D", "produção sob demanda", t)
    t = t.replace("nas variações do anúncio", "nas opções do produto").replace("opções do anúncio", "opções do produto").replace("variação do anúncio", "variação do produto").replace("do anúncio", "do produto").replace("escolhida no anúncio", "escolhida")
    return t


def nome_limpo(n: str) -> str:
    n = n.split(" | ")[0].replace("  ", " ").strip()
    if n.isupper():
        n = " ".join(w.capitalize() if w.lower() not in ("e", "de", "da", "do") else w.lower() for w in n.split())
        n = n[0].upper() + n[1:]
        n = re.sub(r"([“\"])([a-zà-ú])", lambda m: m.group(1) + m.group(2).upper(), n)
    return n


def main(entrada, saida):
    produtos = json.load(open(entrada, encoding="utf-8"))["produtos"]
    shutil.copy("docs/catalogo/modelo-catalogo-afeturar.xlsx", saida)
    wb = openpyxl.load_workbook(saida)
    ws = wb["Produtos"]
    cab = {str(c.value).split(" (")[0].split(" —")[0].strip().lower(): i + 1 for i, c in enumerate(ws[1]) if c.value}
    col = lambda p: next(i for k, i in cab.items() if k.startswith(p))
    for r in range(ws.max_row, 1, -1):
        ws.delete_rows(r)
    linha = 2
    avisos = []
    for p in produtos:
        base = p["preco_min"]
        eixos = [v["nome"] for v in p["variacoes"]]
        cat = CATEGORIAS[p["slug"]]
        desc = sem_3d(p["descricao"])
        if re.search(r"3D", desc):
            avisos.append(f"{p['nome']}: ainda cita 3D")
        dim = p["dimensoes_cm"]
        for i, m in enumerate(p["modelos"]):
            cor, outros = "", []
            for eixo, valor in zip(eixos, m["combinacao"]):
                if not valor:
                    continue
                if eixo == "Cor":
                    cor = corrigir(valor)
                else:
                    outros.append(corrigir(valor))
            dados = {}
            if i == 0:
                dados = {
                    "incluir": "Sim", "produto": nome_limpo(p["nome"]), "categoria": cat, "disponibilidade": "Sob encomenda",
                    "preço de venda": base, "descrição": desc,
                    "embalagem: peso": round(p["peso_kg"] * 1000), "embalagem: comprimento": round(float(dim["length"]) * 10),
                    "embalagem: largura": round(float(dim["width"]) * 10), "embalagem: altura": round(float(dim["height"]) * 10),
                    "observações": "Importado da Shopee. Prazo de preparo e medidas da peça: preencher.",
                }
            else:
                dados = {"produto": nome_limpo(p["nome"])}
            if m["preco"] != base:
                dados["preço da variação"] = m["preco"]
            if cor:
                dados["cor"] = cor
            if outros:
                dados["tamanho"] = " / ".join(outros)
            for k, v in dados.items():
                ws.cell(linha, col(k)).value = v
            linha += 1
    wb.save(saida)
    print(f"{len(produtos)} produtos, {linha - 2} linhas -> {saida}")
    for a in avisos:
        print("AVISO", a)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
